import json
import pathlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_spreadsheets(json_path: pathlib.Path, quiet: bool = False):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    
    base_dir = json_path.parent
    base_name = json_path.stem
    excel_path = base_dir / f"{base_name}.xlsx"
    csv_path = base_dir / f"{base_name}-legivel.csv"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ranking de Ações"
    ws.views.sheetView[0].showGridLines = True

    # Definição das colunas
    headers = [
        ("Pos.", 6),
        ("Ticker", 10),
        ("Recomendação", 16),
        ("Score Ponderado", 16),
        ("Confiança", 12),
        ("P(Compra)", 12),
        ("P(Hold)", 12),
        ("P(Venda)", 12),
        ("Técnica (0-4)", 14),
        ("Fundamentos (0-4)", 18),
        ("Sentimento (0-4)", 16),
        ("Timing (0-4)", 14),
        ("Risco (0-3)", 13),
        ("Status Gates", 14),
        ("Justificativa", 55),
        ("Modelo", 12),
        ("Data/Hora", 20),
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    for col_idx, (header_text, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 26

    # Estilos de recomendação
    compra_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    compra_font = Font(name="Calibri", size=11, bold=True, color="375623")

    hold_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    hold_font = Font(name="Calibri", size=11, bold=True, color="7F6000")

    venda_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    venda_font = Font(name="Calibri", size=11, bold=True, color="C65911")

    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    csv_rows = []
    csv_headers = [h[0] for h in headers]
    csv_rows.append(csv_headers)

    for row_idx, item in enumerate(data, 2):
        ws.row_dimensions[row_idx].height = 22
        is_zebra = (row_idx % 2 == 1)
        row_fill = zebra_fill if is_zebra else None

        pos = row_idx - 1
        ticker = item.get("ticker", "")
        rec = (item.get("recomendacao") or "").upper()
        weighted = item.get("weighted", 0.0)
        confidence = item.get("confidence", 0.0)
        
        probs = item.get("probabilidades", {})
        p_compra = probs.get("compra", 0.0)
        p_hold = probs.get("hold", 0.0)
        p_venda = probs.get("venda", 0.0)

        scores = item.get("scores", {})
        tec = scores.get("tendencia_tecnica", 0.0)
        fund = scores.get("qualidade_fundamentalista", 0.0)
        sent = scores.get("sentimento_noticia", 0.0)
        tim = scores.get("timing_momentum", 0.0)
        risc = scores.get("risco_volatilidade", 0.0)

        motivo_gate = item.get("motivo_gate")
        gate_status = f"Bloqueio: {motivo_gate}" if motivo_gate else "Aprovado"

        just = item.get("justificativa", "")
        model = item.get("model", "")
        ts = item.get("timestamp", "")
        if "T" in ts:
            date_part, time_part = ts.split("T")
            time_clean = time_part.split(".")[0]
            ts_formatted = f"{date_part} {time_clean}"
        else:
            ts_formatted = ts

        row_data = [
            (pos, center_align, "0"),
            (ticker, center_align, None),
            (rec, center_align, None),
            (weighted, right_align, "0.000"),
            (confidence, right_align, "0.0%"),
            (p_compra, right_align, "0.0%"),
            (p_hold, right_align, "0.0%"),
            (p_venda, right_align, "0.0%"),
            (tec, right_align, "0.0"),
            (fund, right_align, "0.0"),
            (sent, right_align, "0.0"),
            (tim, right_align, "0.0"),
            (risc, right_align, "0.0"),
            (gate_status, center_align, None),
            (just, left_align, None),
            (model, center_align, None),
            (ts_formatted, center_align, None),
        ]

        csv_row = [
            str(pos),
            ticker,
            rec,
            f"{weighted:.3f}".replace(".", ","),
            f"{confidence*100:.1f}%".replace(".", ","),
            f"{p_compra*100:.1f}%".replace(".", ","),
            f"{p_hold*100:.1f}%".replace(".", ","),
            f"{p_venda*100:.1f}%".replace(".", ","),
            f"{tec:.1f}".replace(".", ","),
            f"{fund:.1f}".replace(".", ","),
            f"{sent:.1f}".replace(".", ","),
            f"{tim:.1f}".replace(".", ","),
            f"{risc:.1f}".replace(".", ","),
            gate_status,
            f'"{just}"',
            model,
            ts_formatted
        ]
        csv_rows.append(csv_row)

        for col_idx, (val, align, num_fmt) in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = align
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if num_fmt:
                cell.number_format = num_fmt

            # Custom styling
            if col_idx == 2:
                cell.font = Font(name="Calibri", size=11, bold=True)
            elif col_idx == 3:
                if rec == "COMPRA":
                    cell.fill = compra_fill
                    cell.font = compra_font
                elif rec == "VENDA":
                    cell.fill = venda_fill
                    cell.font = venda_font
                else:
                    cell.fill = hold_fill
                    cell.font = hold_font

    # Filtro automático
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"
    # Congelar painel superior
    ws.freeze_panes = "A2"

    wb.save(excel_path)
    if not quiet:
        print(f"  ✓ Salvo Excel: {excel_path}")

    # Salvar cópia latest se for datado
    if "ranking-" in base_name and base_name != "ranking-latest":
        latest_excel = base_dir / "ranking-latest.xlsx"
        wb.save(latest_excel)
        if not quiet:
            print(f"  ✓ Salvo Excel Latest: {latest_excel}")

    # CSV formatado em UTF-8 com BOM e separador ';' para compatibilidade nativa com Excel PT-BR
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        for r in csv_rows:
            f.write(";".join(r) + "\n")
    if not quiet:
        print(f"  ✓ Salvo CSV Legível: {csv_path}")

    if "ranking-" in base_name and base_name != "ranking-latest":
        latest_csv = base_dir / "ranking-latest-legivel.csv"
        with open(latest_csv, "w", encoding="utf-8-sig", newline="") as f:
            for r in csv_rows:
                f.write(";".join(r) + "\n")
        if not quiet:
            print(f"  ✓ Salvo CSV Latest Legível: {latest_csv}")


if __name__ == "__main__":
    import sys
    target = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path("/Users/jpoloni/dev/jev/data/ranking-2026-09-20.json")
    generate_spreadsheets(target)
